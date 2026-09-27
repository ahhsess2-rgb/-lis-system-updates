# -*- coding: utf-8 -*-
"""
excel_export.py
=============================================================================
تصدير Excel (.xlsx) عام لأي جدول نتائج — يُستخدم أول شي من ميزة "بحث
وتصدير الحسابات" الجديدة، ويصلح لأي تصدير جدولي مستقبلي بنفس الفكرة.

يحتاج مكتبة openpyxl (غير مثبتة سابقًا بهذا المشروع حسب فحصي):
    pip install openpyxl

الاستخدام:
    import excel_export
    path = excel_export.export_rows_to_xlsx(
        columns=[("test_name", "التحليل"), ("date", "التاريخ"), ("price", "السعر")],
        rows=[{"test_name": "CBC", "date": "12/09/2026", "price": 5000}, ...],
        sheet_title="نتائج البحث",
        totals_row={"price": 150000},     # اختياري: صف مجموع بالأسفل
        rtl=True,
    )
    # path هو مسار ملف مؤقت جاهز لـ send_file(path, as_attachment=True, download_name=...)
"""
import os
import tempfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def make_temp_xlsx_path(prefix):
    """نفس فكرة pdf_export.make_temp_pdf_path الموجودة أصلاً بالمشروع —
    مسار ملف مؤقت فريد بامتداد xlsx."""
    ts = datetime.now().strftime("%Y%m%d%H%M%S")
    fd, path = tempfile.mkstemp(prefix=f"{prefix}_{ts}_", suffix=".xlsx")
    os.close(fd)
    return path


def export_rows_to_xlsx(columns, rows, sheet_title="Sheet1", totals_row=None,
                         rtl=True, out_path=None):
    """columns: [(key, header_label), ...] — بنفس ترتيب الأعمدة المطلوب.
    rows: [{key: value, ...}, ...].
    totals_row: dict اختياري {key: value} يُكتب كصف أخير بخط عريض (مثلاً
                {"price": 150000, "test_name": "المجموع"}).
    rtl: يجعل اتجاه الورقة RTL (مناسب لعناوين الأعمدة العربية).
    out_path: مسار جاهز لو ما تحتاج ملف مؤقت جديد (يُستخدم make_temp_xlsx_path لو فاضي).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Sheet1")[:31]  # Excel sheet name hard limit
    if rtl:
        ws.sheet_view.rightToLeft = True

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="205072", end_color="205072", fill_type="solid")
    totals_font = Font(bold=True)

    for col_idx, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    if totals_row:
        r = len(rows) + 2
        for col_idx, (key, _) in enumerate(columns, start=1):
            if key in totals_row:
                cell = ws.cell(row=r, column=col_idx, value=totals_row[key])
                cell.font = totals_font

    # عرض أعمدة تلقائي تقريبي حسب أطول محتوى بكل عمود (Excel ما عنده
    # auto-fit حقيقي عبر openpyxl، هذا تقريب كافٍ عمليًا).
    for col_idx, (key, label) in enumerate(columns, start=1):
        max_len = len(str(label))
        for row in rows:
            val = row.get(key, "")
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    path = out_path or make_temp_xlsx_path("accounts_search")
    wb.save(path)
    return path
